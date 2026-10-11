"""Extractor panel — the different ways real ATS stacks turn a CV file into text.

An ATS never "reads the PDF"; it runs ONE text extractor and parses whatever comes out.
Different stacks use different extractors, and they disagree on reading order (columns,
side-by-side meta), on what they skip (headers, text boxes) and on Unicode mapping. A CV is
only robust if it survives ALL of them, so the lab runs every file through the whole panel.

PDF : poppler pdftotext (raw + -layout), pdfminer.six, pypdf, pdfplumber, PyMuPDF, pdf.js
DOCX: python-docx, docx2txt, PyMuPDF, LibreOffice, mammoth
+ Apache Tika (PDFBox/POI — the common Java enterprise stack) when TIKA_JAR is set (CI).
"""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Callable

LAB_NODE = Path(__file__).resolve().parents[3] / "web" / "lab"  # apps/web/lab


def _run(cmd: list[str], timeout: int = 60) -> str:
    out = subprocess.run(cmd, capture_output=True, timeout=timeout, check=True)
    return out.stdout.decode("utf-8", errors="replace")


# ----------------------------------------------------------------------------- PDF
def _pdftotext(path: str) -> str:
    return _run(["pdftotext", "-enc", "UTF-8", path, "-"])


def _pdftotext_layout(path: str) -> str:
    return _run(["pdftotext", "-layout", "-enc", "UTF-8", path, "-"])


def _pdfminer(path: str) -> str:
    from pdfminer.high_level import extract_text

    return extract_text(path)


def _pypdf(path: str) -> str:
    from pypdf import PdfReader

    return "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)


def _pdfplumber(path: str) -> str:
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        return "\n".join((p.extract_text() or "") for p in pdf.pages)


def _pymupdf(path: str) -> str:
    import pymupdf

    with pymupdf.open(path) as doc:
        return "\n".join(page.get_text() for page in doc)


def _pdfjs(path: str) -> str:
    # pdf.js text items joined into lines by y — what browser-side parsers (OpenResume,
    # ats-screener, our own client analyzer) see.
    return _run(["node", str(LAB_NODE / "pdfjs-extract.mjs"), path])


# ---------------------------------------------------------------------------- DOCX
def _python_docx(path: str) -> str:
    import docx

    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)


def _docx2txt(path: str) -> str:
    import docx2txt

    return docx2txt.process(path)


def _pymupdf_docx(path: str) -> str:
    import pymupdf

    with pymupdf.open(path, filetype="docx") as doc:
        return "\n".join(page.get_text() for page in doc)


def _libreoffice(path: str) -> str:
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "txt:Text (encoded):UTF8", "--outdir", tmp, path],
            capture_output=True, timeout=120, check=True,
        )
        out = Path(tmp) / (Path(path).stem + ".txt")
        return out.read_text(encoding="utf-8", errors="replace")


def _mammoth(path: str) -> str:
    return _run(["node", str(LAB_NODE / "mammoth-extract.mjs"), path])


def _tika(path: str) -> str:
    # Apache Tika (PDFBox for PDF, POI for DOCX): the common Java enterprise ATS stack.
    # Enabled when TIKA_JAR points at tika-app-<version>.jar (CI downloads it).
    return _run(["java", "-jar", os.environ["TIKA_JAR"], "--text", path], timeout=180)


PDF_EXTRACTORS: dict[str, Callable[[str], str]] = {
    "pdftotext": _pdftotext,
    "pdftotext-layout": _pdftotext_layout,
    "pdfminer": _pdfminer,
    "pypdf": _pypdf,
    "pdfplumber": _pdfplumber,
    "pymupdf": _pymupdf,
    "pdfjs": _pdfjs,
    "tika": _tika,
}

DOCX_EXTRACTORS: dict[str, Callable[[str], str]] = {
    "python-docx": _python_docx,
    "docx2txt": _docx2txt,
    "pymupdf": _pymupdf_docx,
    "libreoffice": _libreoffice,
    "mammoth": _mammoth,
    "tika": _tika,
}


def available(extractors: dict[str, Callable[[str], str]]) -> dict[str, Callable[[str], str]]:
    """Drop extractors whose binary/runtime is missing on this machine (CI vs laptop)."""
    need = {"pdftotext": "pdftotext", "pdftotext-layout": "pdftotext", "pdfjs": "node",
            "libreoffice": "soffice", "mammoth": "node"}
    present = {k: f for k, f in extractors.items() if k not in need or shutil.which(need[k])}
    # Node readers also need their package installed (`npm ci` in apps/web); a missing local
    # dependency is an environment gap, not a CV failure, so it must not reach the gate.
    node_modules = LAB_NODE.parent / "node_modules"
    for key, pkg in (("pdfjs", "pdfjs-dist"), ("mammoth", "mammoth")):
        if key in present and not (node_modules / pkg).is_dir():
            present.pop(key)
    if not (os.environ.get("TIKA_JAR") and shutil.which("java")):
        present.pop("tika", None)
    return present


def extract_all(path: str) -> dict[str, str | Exception]:
    """Run the whole panel; an extractor crash is a result too (the ATS would get nothing)."""
    panel = PDF_EXTRACTORS if path.lower().endswith(".pdf") else DOCX_EXTRACTORS
    results: dict[str, str | Exception] = {}
    for name, fn in available(panel).items():
        try:
            results[name] = fn(path)
        except Exception as exc:  # noqa: BLE001 — recorded, not raised
            results[name] = exc
    return results


if __name__ == "__main__":  # quick manual check: python -m app.lab.extractors file.pdf
    import sys

    for name, text in extract_all(sys.argv[1]).items():
        print(f"===== {name} =====")
        print(text if isinstance(text, str) else f"ERROR: {text}")
